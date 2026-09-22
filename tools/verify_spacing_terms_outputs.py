"""Final structural checks for the corrected HandVox chapters 1–3."""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "ผลลัพธ์เอกสาร"
EXPECTED_STARTS = {1: 1, 2: 6, 3: 37}
EXPECTED_COUNTS = {
    1: (48, 1, 0),
    2: (284, 9, 0),
    3: (67, 3, 2),
}
FONT = "TH Sarabun New"


def all_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        for part in (
            section.header,
            section.first_page_header,
            section.even_page_header,
            section.footer,
            section.first_page_footer,
            section.even_page_footer,
        ):
            yield from part.paragraphs


def main():
    failed = False
    for chapter in range(1, 4):
        path = OUTPUT / f"บทที่ {chapter}.docx"
        document = Document(path)
        section = document.sections[0]
        pg_num_type = section._sectPr.find(qn("w:pgNumType"))
        page_start = int(pg_num_type.get(qn("w:start")))
        header_xml = section.header._element.xml
        has_page_field = "PAGE" in header_xml
        bad_fonts = []
        missing_language = 0
        has_word_joiner = False
        for paragraph in all_paragraphs(document):
            has_word_joiner = has_word_joiner or "\u2060" in paragraph.text
            for run in paragraph.runs:
                if not run.text:
                    continue
                rpr = run._r.get_or_add_rPr()
                fonts = rpr.rFonts
                if fonts is not None:
                    names = {
                        fonts.get(qn("w:ascii")),
                        fonts.get(qn("w:hAnsi")),
                        fonts.get(qn("w:eastAsia")),
                        fonts.get(qn("w:cs")),
                    } - {None}
                    if any(name != FONT for name in names):
                        bad_fonts.extend(sorted(names - {FONT}))
                lang = rpr.find(qn("w:lang"))
                if lang is None or lang.get(qn("w:eastAsia")) != "th-TH":
                    missing_language += 1

        actual_counts = (
            len(document.paragraphs),
            len(document.tables),
            len(document.inline_shapes),
        )
        ok = (
            page_start == EXPECTED_STARTS[chapter]
            and section.different_first_page_header_footer
            and has_page_field
            and actual_counts == EXPECTED_COUNTS[chapter]
            and not bad_fonts
            and missing_language == 0
            and not has_word_joiner
        )
        failed = failed or not ok
        print(
            f"chapter={chapter} ok={ok} page_start={page_start} "
            f"first_page_hidden={section.different_first_page_header_footer} "
            f"page_field={has_page_field} counts={actual_counts} "
            f"bad_fonts={len(bad_fonts)} missing_language={missing_language} "
            f"word_joiner={has_word_joiner} size={path.stat().st_size}"
        )
    if failed:
        raise SystemExit(1)


if __name__ == "__main__":
    main()

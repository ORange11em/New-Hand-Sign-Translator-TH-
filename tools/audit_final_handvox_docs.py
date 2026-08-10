"""ตรวจเอกสาร Word ฉบับสุดท้ายว่ามีหัวข้อ รูป ตาราง และข้อความสำคัญครบ."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "ผลลัพธ์เอกสาร"
EXPECTED_PLACEHOLDERS = {1: 0, 2: 2, 3: 2, 4: 2, 5: 0}
EXPECTED_PAGE_STARTS = {1: 1, 2: 8, 3: 17, 4: 26, 5: 34}
STALE_TERMS = ("4,500", "15 ป้ายกำกับ", "90 คลิป", "99.11", "90.78")


def document_text(document: Document) -> str:
    chunks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def main() -> None:
    failed = False
    for chapter in range(1, 6):
        path = OUTPUT / f"บทที่ {chapter}_ฉบับปรับปรุง_HandVox.docx"
        document = Document(path)
        text = document_text(document)
        placeholder_count = text.count("[เว้นพื้นที่สำหรับภาพประกอบ]")
        stale = [term for term in STALE_TERMS if term in text]
        section = document.sections[0]
        a4 = abs(section.page_width.mm - 210) < 1 and abs(section.page_height.mm - 297) < 1
        pg_num_type = section._sectPr.find(qn("w:pgNumType"))
        page_start = int(pg_num_type.get(qn("w:start"))) if pg_num_type is not None else None
        first_page_hidden = section.different_first_page_header_footer and not section.first_page_header.paragraphs[0].text.strip()
        page_field_top_right = (
            section.header.paragraphs[0].alignment == 2
            and "PAGE" in "".join(section.header.paragraphs[0]._p.itertext())
        )
        heading_indents_ok = True
        heading_counts = {1: 0, 2: 0, 3: 0}
        nonblack_runs = 0
        for paragraph in document.paragraphs:
            if paragraph.style and paragraph.style.name.startswith("Heading "):
                level = int(paragraph.style.name.split()[-1])
                heading_counts[level] += 1
                actual = paragraph.paragraph_format.left_indent or 0
                expected = Cm(1.25 * (level - 1))
                heading_indents_ok = heading_indents_ok and abs(int(actual) - int(expected)) <= 500
            for run in paragraph.runs:
                color = run.font.color.rgb
                if color is not None and str(color) != "000000":
                    nonblack_runs += 1
        body_children = list(document.element.body)
        table_captions = []
        captions_below = True
        for index, child in enumerate(body_children):
            if child.tag != qn("w:p"):
                continue
            child_text = "".join(child.itertext()).strip()
            if not child_text.startswith("ตารางที่ "):
                continue
            table_captions.append(child_text)
            captions_below = captions_below and index > 0 and body_children[index - 1].tag == qn("w:tbl")
        ok = (
            placeholder_count == EXPECTED_PLACEHOLDERS[chapter]
            and not stale
            and a4
            and len(table_captions) == len(document.tables)
            and captions_below
            and page_start == EXPECTED_PAGE_STARTS[chapter]
            and first_page_hidden
            and page_field_top_right
            and heading_indents_ok
            and nonblack_runs == 0
        )
        failed = failed or not ok
        print(
            f"chapter={chapter} ok={ok} placeholders={placeholder_count} "
            f"tables={len(document.tables)} paragraphs={len(document.paragraphs)} "
            f"table_captions={len(table_captions)} captions_below={captions_below} "
            f"a4={a4} page_start={page_start} first_hidden={first_page_hidden} "
            f"page_top_right={page_field_top_right} headings={heading_counts} "
            f"heading_indents={heading_indents_ok} nonblack_runs={nonblack_runs} stale={stale}"
        )
    if failed:
        raise SystemExit(1)


if __name__ == "__main__":
    main()

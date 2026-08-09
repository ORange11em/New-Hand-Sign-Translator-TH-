"""อ่านโครงสร้าง DOCX เพื่อสรุปย่อหน้า ตาราง รูป และ style สำหรับตรวจเอกสาร."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document


def inspect(path: Path) -> dict:
    doc = Document(path)
    paragraphs = []
    for index, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text:
            paragraphs.append({
                "index": index,
                "style": p.style.name if p.style else "",
                "text": text,
            })

    tables = []
    for table_index, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append({"index": table_index, "rows": rows})

    sections = []
    for s in doc.sections:
        sections.append({
            "page_width": s.page_width,
            "page_height": s.page_height,
            "top_margin": s.top_margin,
            "bottom_margin": s.bottom_margin,
            "left_margin": s.left_margin,
            "right_margin": s.right_margin,
        })

    return {
        "path": str(path),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "inline_shape_count": len(doc.inline_shapes),
        "sections": sections,
        "paragraphs": paragraphs,
        "tables": tables,
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("path", type=Path)
    parser.add_argument("--start", type=int)
    parser.add_argument("--end", type=int)
    parser.add_argument("--no-tables", action="store_true")
    args = parser.parse_args()
    result = inspect(args.path)
    if args.start is not None or args.end is not None:
        start = 0 if args.start is None else args.start
        end = 10**9 if args.end is None else args.end
        result["paragraphs"] = [
            p for p in result["paragraphs"] if start <= p["index"] <= end
        ]
    if args.no_tables:
        result["tables"] = []
    print(json.dumps(result, ensure_ascii=False, indent=2))

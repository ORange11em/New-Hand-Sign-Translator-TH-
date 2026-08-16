"""สร้างรายงาน HandVox 5 บทฉบับรวม พร้อมเลขหน้าต่อเนื่องทุกบท."""

from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT / "tools") not in sys.path:
    sys.path.insert(0, str(ROOT / "tools"))

import build_handvox_chapters as chapters


OUTPUT = ROOT / "ผลลัพธ์เอกสาร" / "HandVox_รายงานฉบับสมบูรณ์_5บท.docx"


def _copy_body_with_images(target, source):
    """คัดลอกเนื้อหา source และสร้าง relationship ของรูปใน target ใหม่."""
    target_body = target.element.body
    final_sect_pr = target_body.sectPr
    insert_at = target_body.index(final_sect_pr)

    for child in source.element.body.iterchildren():
        if child.tag == qn("w:sectPr"):
            continue
        cloned = deepcopy(child)
        for blip in cloned.xpath(".//a:blip"):
            old_rid = blip.get(qn("r:embed"))
            if not old_rid:
                continue
            related_part = source.part.related_parts.get(old_rid)
            if related_part is None or not hasattr(related_part, "blob"):
                continue
            new_rid, _ = target.part.get_or_add_image(BytesIO(related_part.blob))
            blip.set(qn("r:embed"), new_rid)
        target_body.insert(insert_at, cloned)
        insert_at += 1


def _configure_combined_section(section, first=False):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.81)
    section.left_margin = Cm(3.81)
    section.right_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)
    section.different_first_page_header_footer = True

    pg_num_type = section._sectPr.find(qn("w:pgNumType"))
    if pg_num_type is None and first:
        pg_num_type = OxmlElement("w:pgNumType")
        section._sectPr.append(pg_num_type)
    if pg_num_type is not None:
        if first:
            pg_num_type.set(qn("w:start"), "1")
        elif qn("w:start") in pg_num_type.attrib:
            del pg_num_type.attrib[qn("w:start")]

    section.header.is_linked_to_previous = False
    regular_header = section.header.paragraphs[0]
    regular_header.clear()
    chapters.add_page_number(regular_header)

    section.first_page_header.is_linked_to_previous = False
    section.first_page_header.paragraphs[0].clear()
    section.footer.is_linked_to_previous = False
    section.footer.paragraphs[0].clear()
    section.first_page_footer.is_linked_to_previous = False
    section.first_page_footer.paragraphs[0].clear()


def build():
    chapter_paths = [
        chapters.build_chapter_1(),
        chapters.build_chapter_2(),
        chapters.build_chapter_3(),
        chapters.build_chapter_4(),
        chapters.build_chapter_5(),
    ]

    combined = Document(chapter_paths[0])
    _configure_combined_section(combined.sections[0], first=True)

    for path in chapter_paths[1:]:
        new_section = combined.add_section(WD_SECTION.NEW_PAGE)
        _configure_combined_section(new_section, first=False)
        source = Document(path)
        _copy_body_with_images(combined, source)

    settings = combined.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    combined.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())

"""ตรวจรูปแบบสำคัญของรายงาน HandVox ฉบับรวมก่อนส่งมอบ."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


FONT = "TH Sarabun New"
BLACK = "000000"


def iter_blocks(document):
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def all_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        for story in (section.header, section.first_page_header, section.footer, section.first_page_footer):
            yield from story.paragraphs


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    args = parser.parse_args()
    doc = Document(args.docx)
    errors = []

    if len(doc.sections) != 5:
        errors.append(f"sections={len(doc.sections)} (expected 5)")

    for index, section in enumerate(doc.sections, start=1):
        if round(section.page_width.cm, 1) != 21.0 or round(section.page_height.cm, 1) != 29.7:
            errors.append(f"section {index}: not A4")
        margins = tuple(round(value.cm, 2) for value in (section.top_margin, section.left_margin, section.right_margin, section.bottom_margin))
        if margins != (3.81, 3.81, 2.54, 2.54):
            errors.append(f"section {index}: margins={margins}")
        if not section.different_first_page_header_footer:
            errors.append(f"section {index}: first-page header is not different")
        regular_xml = section.header._element.xml
        first_xml = section.first_page_header._element.xml
        if " PAGE " not in regular_xml:
            errors.append(f"section {index}: missing PAGE field in regular header")
        if " PAGE " in first_xml:
            errors.append(f"section {index}: PAGE field appears on chapter-opening page")
        pg_num_type = section._sectPr.find(qn("w:pgNumType"))
        start = None if pg_num_type is None else pg_num_type.get(qn("w:start"))
        if index == 1 and start != "1":
            errors.append(f"section 1: page start={start!r}, expected '1'")
        if index > 1 and start is not None:
            errors.append(f"section {index}: page numbering restarts at {start}")

    for paragraph in all_paragraphs(doc):
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            if run.font.name not in (FONT, None):
                errors.append(f"wrong font {run.font.name!r}: {run.text[:40]!r}")
            rgb = run.font.color.rgb
            if rgb is not None and str(rgb).upper() != BLACK:
                errors.append(f"non-black text {rgb}: {run.text[:40]!r}")

    expected_indents = {"Heading 1": 0.0, "Heading 2": 1.25, "Heading 3": 2.5}
    for paragraph in doc.paragraphs:
        if paragraph.style.name not in expected_indents:
            continue
        actual = round((paragraph.paragraph_format.left_indent.cm if paragraph.paragraph_format.left_indent else 0.0), 2)
        if actual != expected_indents[paragraph.style.name]:
            errors.append(f"{paragraph.style.name} indent={actual}: {paragraph.text[:50]!r}")

    blocks = list(iter_blocks(doc))
    table_count = 0
    for idx, block in enumerate(blocks):
        if not isinstance(block, Table):
            continue
        table_count += 1
        if idx + 1 >= len(blocks) or not isinstance(blocks[idx + 1], Paragraph) or not blocks[idx + 1].text.strip().startswith("ตารางที่"):
            errors.append(f"table {table_count}: caption is not immediately below table")
    if table_count != 22:
        errors.append(f"tables={table_count} (expected 22)")

    body_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    for chapter in range(1, 6):
        if f"บทที่ {chapter}" not in body_text:
            errors.append(f"missing chapter title {chapter}")
    for stale in ("ข้อมูล 120 คลิปจาก 4 ท่า", "accuracy 100%", "ชุดคำศัพท์ที่เสนอให้เพิ่มในอนาคต 16 ท่า"):
        if stale.lower() in body_text.lower():
            errors.append(f"stale text remains: {stale}")
    if body_text.count("[เว้นพื้นที่สำหรับภาพประกอบ]") != 3:
        errors.append("expected 3 labeled image placeholders")

    if len(doc.inline_shapes) != 3:
        errors.append(f"inline images={len(doc.inline_shapes)} (expected 3)")
    for index, shape in enumerate(doc.inline_shapes, start=1):
        descr = shape._inline.docPr.get("descr")
        if not descr:
            errors.append(f"image {index}: missing alt text")

    if errors:
        print("FAIL")
        for error in errors:
            print(f"- {error}")
        raise SystemExit(1)
    print(f"PASS: 5 sections, {table_count} tables, {len(doc.inline_shapes)} images, continuous page numbering, TH Sarabun New, black text")


if __name__ == "__main__":
    main()

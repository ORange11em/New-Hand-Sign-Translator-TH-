from pathlib import Path
from copy import deepcopy
import json, re, subprocess
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'ผลลัพธ์เอกสาร'
WORK=Path(__file__).parent
FONT='TH Sarabun New'

def child(parent,name,attrs=None):
    e=parent.find(qn('w:'+name))
    if e is None:
        e=OxmlElement('w:'+name); parent.append(e)
    for k,v in (attrs or {}).items(): e.set(qn('w:'+k),str(v))
    return e

def font(run,size=16,bold=False):
    run.font.name=FONT; run.font.size=Pt(size); run.font.bold=bold
    run.font.color.rgb=RGBColor(0,0,0)
    rp=run._r.get_or_add_rPr()
    rf=child(rp,'rFonts')
    for k in ['ascii','hAnsi','eastAsia','cs']: rf.set(qn('w:'+k),FONT)
    for k in list(rf.attrib):
        if 'theme' in k.lower(): del rf.attrib[k]
    child(rp,'lang',{'val':'th-TH','eastAsia':'th-TH','bidi':'th-TH'})
    child(rp,'szCs',{'val':str(size*2)})
    child(rp,'bCs',{'val':'1' if bold else '0'})
    child(rp,'spacing',{'val':'0'})
    child(rp,'w',{'val':'100'})

def setup(doc,start):
    sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(3.81);sec.bottom_margin=Cm(2.54)
    sec.left_margin=Cm(3.81);sec.right_margin=Cm(2.54)
    sec.header_distance=Cm(1.5);sec.footer_distance=Cm(1.25)
    sec.different_first_page_header_footer=True
    child(sec._sectPr,'pgNumType',{'start':start})
    for hf in [sec.header,sec.footer,sec.first_page_header,sec.first_page_footer,sec.even_page_header,sec.even_page_footer]:
        for p in list(hf.paragraphs): p.clear()
    hp=sec.header.paragraphs[0];hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_before=Pt(0);hp.paragraph_format.space_after=Pt(0)
    r=hp.add_run();font(r,16)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),' PAGE ')
    rr=OxmlElement('w:r');rp=OxmlElement('w:rPr')
    rf=OxmlElement('w:rFonts')
    for k in ['ascii','hAnsi','cs','eastAsia']:rf.set(qn('w:'+k),FONT)
    rp.append(rf);child(rp,'sz',{'val':'32'});child(rp,'szCs',{'val':'32'})
    rr.append(rp);t=OxmlElement('w:t');t.text=str(start+1);rr.append(t);fld.append(rr);hp._p.append(fld)
    for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
        st=doc.styles[name];st.font.name=FONT;st.font.size=Pt(16);st.font.color.rgb=RGBColor(0,0,0)
        st.font.bold=name!='Normal';st.font.italic=False
        pf=st.paragraph_format;pf.space_before=Pt(0);pf.space_after=Pt(4)
        pf.line_spacing=Pt(19);pf.line_spacing_rule=WD_LINE_SPACING.EXACTLY
        pf.keep_together=False;pf.keep_with_next=name!='Normal';pf.widow_control=True
        pf.left_indent=Cm(0);pf.first_line_indent=Cm(0)
        rp=st.element.get_or_add_rPr();rf=child(rp,'rFonts')
        for k in ['ascii','hAnsi','cs','eastAsia']:rf.set(qn('w:'+k),FONT)
        for k in list(rf.attrib):
            if 'theme' in k.lower():del rf.attrib[k]
        child(rp,'szCs',{'val':'32'});child(rp,'lang',{'val':'th-TH','bidi':'th-TH','eastAsia':'th-TH'})
        for n in rp.findall(qn('w:color')):
            for a in list(n.attrib):
                if 'theme' in a.lower():del n.attrib[a]
    settings=doc.settings.element
    for name in ['evenAndOddHeaders','trackRevisions','updateFields']:
        for e in list(settings.findall(qn('w:'+name))):settings.remove(e)
    child(settings,'updateFields',{'val':'true'})
    child(settings,'characterSpacingControl',{'val':'doNotCompress'})
    comp=child(settings,'compat');child(comp,'doNotExpandShiftReturn')
    for setting in comp.findall(qn('w:compatSetting')):
        if setting.get(qn('w:name'))=='compatibilityMode':setting.set(qn('w:val'),'15')
    defaults=doc.styles.element.find(qn('w:docDefaults'))
    if defaults is not None:
        for rf in defaults.findall('.//'+qn('w:rFonts')):
            rf.attrib.clear()
            for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),FONT)
    # Remove inherited title rules and character-level wrapping from all styles.
    for tag in ['pBdr','wordWrap']:
        for e in list(doc.styles.element.iter(qn('w:'+tag))):e.getparent().remove(e)

def para(doc,text,kind='body',level=0):
    if kind=='body':
        text=re.sub(r' (\[\d+\](?:, \[\d+\])*)$',lambda m:'\u00a0'+m.group(1).replace(' ','\u00a0'),text)
    if kind=='body' and re.match(r'^\d\)',text):
        number, heading, body=text.split(' ',2)
        p=doc.add_paragraph(style='Normal')
        pf=p.paragraph_format
        pf.left_indent=Cm(1.6);pf.first_line_indent=Cm(-.6)
        pf.tab_stops.add_tab_stop(Cm(1.6))
        pf.space_before=Pt(3);pf.space_after=Pt(0)
        pf.keep_with_next=True;pf.keep_together=True
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        font(p.add_run(number+'\t'+heading),16,True)
        p=para(doc,body,'body',level)
        p.paragraph_format.left_indent=Cm(1.6)
        p.paragraph_format.first_line_indent=Cm(0)
        return p
    p=doc.add_paragraph(style={'body':'Normal','title':'Title','h1':'Heading 1','h2':'Heading 2'}[kind])
    pf=p.paragraph_format
    pf.left_indent=Cm(level*.5);pf.right_indent=Cm(0);pf.first_line_indent=Cm(.75 if kind=='body' else 0)
    pf.space_before=Pt(6 if kind=='h1' else (4 if kind=='h2' else 0));pf.space_after=Pt(4)
    pf.keep_with_next=kind!='body';pf.keep_together=kind!='body';pf.widow_control=True
    pf.page_break_before=False
    p.alignment=WD_ALIGN_PARAGRAPH.THAI_JUSTIFY if kind=='body' else WD_ALIGN_PARAGRAPH.LEFT
    size=16
    if kind=='title':
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.line_spacing=Pt(26);size=20
        pf.first_line_indent=Cm(0);pf.left_indent=Cm(0);pf.space_after=Pt(6)
    pr=p._p.get_or_add_pPr()
    child(pr,'snapToGrid',{'val':'0'})
    child(pr,'autoSpaceDE',{'val':'0'});child(pr,'autoSpaceDN',{'val':'0'})
    font(p.add_run(text),size,kind!='body')
    return p

def objectives_table(doc):
    caption=para(doc,'ตารางที่ 5.1 สรุปผลการดำเนินงานตามวัตถุประสงค์','body',1)
    caption.alignment=WD_ALIGN_PARAGRAPH.LEFT
    caption.paragraph_format.first_line_indent=Cm(0)
    caption.paragraph_format.keep_with_next=True
    caption.paragraph_format.page_break_before=True
    for r in caption.runs:font(r,16,True)
    objective_doc=Document(ROOT/'บทที่ 1_เขียนใหม่_2569-09-23'/'บทที่ 1_บทนำ_HandVox_ฉบับขัดเกลาภาษาวิชาการ.docx')
    objectives=[];active=False
    for p in objective_doc.paragraphs:
        text=p.text.replace('\u200b','')
        if text.startswith('1.2 '):active=True;continue
        if text.startswith('1.3 '):break
        if active and text.strip():objectives.append(re.sub(r'^\d+\.\s*','',text))
    assert len(objectives)==4
    evidence=[
        'ออกแบบขั้นตอนรับภาพ เก็บคลิป และแสดงผลเป็นข้อความกับเสียง พร้อมพัฒนาส่วนติดต่อผู้ใช้ [1], [2]',
        'พัฒนาส่วนรู้จำท่า จัดการข้อความ เรียกใช้ระบบเสียง และเตรียมข้อมูลเพื่อฝึกแบบจำลองใหม่ [1], [3]',
        'มีรายงานตัวชี้วัดและเมทริกซ์ความสับสน แยกผู้ทำท่าเดิมกับคนใหม่ โดย Accuracy เท่ากับ 76.15% และ 90.20% ตามลำดับ [3], [6]',
        'มีรายการตรวจสอบ FT-01 ถึง FT-09 และชุดทดสอบอัตโนมัติผ่าน 26 กรณี ส่วนการฟังเสียงจริงยังต้องทดสอบเพิ่มเติม [5]',
    ]
    evidence=[re.sub(r' \[\d+\](?:, \[\d+\])*$', '', value) for value in evidence]
    status=['ดำเนินการแล้ว\nตามขอบเขต','พัฒนาต้นแบบแล้ว\nคุณภาพการรู้จำ\nยังต้องปรับปรุง','ประเมินแล้ว\nผลผู้ทำท่าใหม่\nยังเป็นผลเบื้องต้น','ตรวจสอบแล้ว\nบางส่วน']
    table=doc.add_table(rows=1,cols=3);table.alignment=WD_TABLE_ALIGNMENT.LEFT;table.autofit=False
    widths=[5.2,5.65,3.3]
    pr=table._tbl.tblPr
    child(pr,'tblW',{'w':str(round(Cm(sum(widths)).twips)),'type':'dxa'})
    child(pr,'tblInd',{'w':str(round(Cm(.5).twips)),'type':'dxa'})
    borders=child(pr,'tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        child(borders,side,{'val':'single','sz':'6','color':'D9D9D9'})
    margins=child(pr,'tblCellMar')
    for side,w in [('top',90),('bottom',90),('left',100),('right',100)]:child(margins,side,{'w':w,'type':'dxa'})
    for col,width in zip(table.columns,widths):col.width=Cm(width)
    headers=['วัตถุประสงค์ของโครงงาน','ผลการดำเนินงาน','สถานะ']
    rows=[headers]+[[f'{i+1}. {o}',evidence[i],status[i]] for i,o in enumerate(objectives)]
    node=Path('C:/Users/Ult_Orange/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe')
    result=subprocess.run([str(node),str(WORK/'segment_thai.mjs'),'--stdin'],input=json.dumps([t for row in rows for t in row],ensure_ascii=False),text=True,encoding='utf-8',capture_output=True,check=True)
    segmented=json.loads(result.stdout)
    rows=[segmented[i:i+3] for i in range(0,len(segmented),3)]
    for ri,texts in enumerate(rows):
        row=table.rows[0] if ri==0 else table.add_row()
        trpr=row._tr.get_or_add_trPr();child(trpr,'cantSplit')
        if ri==0:child(trpr,'tblHeader')
        for ci,(cell,text) in enumerate(zip(row.cells,texts)):
            cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri==0:child(cell._tc.get_or_add_tcPr(),'shd',{'fill':'E7E6E6','val':'clear'})
            p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ri==0 or ci==2 else WD_ALIGN_PARAGRAPH.LEFT
            pf=p.paragraph_format;pf.left_indent=Cm(0);pf.first_line_indent=Cm(0);pf.right_indent=Cm(0)
            pf.space_before=Pt(0);pf.space_after=Pt(0);pf.line_spacing=Pt(18)
            pf.keep_with_next=ri==0;pf.keep_together=True
            child(p._p.get_or_add_pPr(),'snapToGrid',{'val':'0'})
            font(p.add_run(text),14,ri==0)
    return table

doc=Document(OUT/'บทที่ 5.docx')
for node in list(doc._element.body):
    if node.tag!=qn('w:sectPr'):doc._element.body.remove(node)
setup(doc,32)
blocks=json.loads((WORK/'segmented.json').read_text(encoding='utf-8'))
for line in blocks[0].splitlines():para(doc,line,'title')
level=0
for text in blocks[1:]:
    raw=text.replace('\u200b','')
    if raw=='[[OBJECTIVES_TABLE]]':
        objectives_table(doc);continue
    if re.match(r'^5\.\d+\.\d+ ',raw):kind='h2';level=2
    elif re.match(r'^5\.\d+ ',raw):kind='h1';level=1
    else:kind='body'
    para(doc,text,kind,level)
doc.core_properties.title='บทที่ 5 สรุปผล อภิปรายผล และข้อเสนอแนะ'
doc.core_properties.subject='โครงงาน HandVox ระบบรู้จำภาษามือไทยเป็นข้อความและเสียงพูด'
doc.core_properties.author='คณะผู้จัดทำโครงงาน HandVox'
doc.save(OUT/'บทที่ 5.docx')

refs=Document();setup(refs,1)
para(refs,'เอกสารอ้างอิงประกอบบทที่ 5','title')
para(refs,'โครงงาน HandVox ระบบรู้จำภาษามือไทยเป็นข้อความและเสียงพูด','body')
para(refs,'รายการแหล่งข้อมูลประกอบบทที่ 5 สำหรับนำไปรวมในเอกสารอ้างอิงท้ายเล่ม','body')
entries=[
('[1] คณะผู้พัฒนา HandVox. (2569). ซอร์สโค้ดแอปพลิเคชัน ระบบรู้จำ การจัดการข้อความ และการสังเคราะห์เสียงพูด [ซอร์สโค้ดภายในโครงงาน].','handvox/ui.py; handvox/detector_view.py; handvox/training_workflow.py; run_detector.py; tts_service.py'),
('[2] คณะผู้พัฒนา HandVox. (2569). ประวัติข้อความและภาพส่วนติดต่อผู้ใช้ของแอปพลิเคชัน HandVox [ข้อมูลและภาพประกอบภายในโครงงาน].','conversation_history.json; ชุดรูปประกอบเอกสาร_01_02_03/03_รูปโปรแกรม; ผลลัพธ์เอกสาร/บทที่ 4.docx'),
('[3] คณะผู้พัฒนา HandVox. (2569, 15 กันยายน). รายงานผลการฝึกและประเมินแบบจำลอง TCN และสถานะแบบจำลองที่ติดตั้ง [รายงานผลการทดลองภายในโครงงาน].','experiments/20260915_061512_537476/report.md; metrics.json; confusion_matrix.csv; model_manifest.json ของแบบจำลองที่ติดตั้ง'),
('[4] คณะผู้พัฒนา HandVox. (2569). บันทึกความเร็วการทำงานของหน้าจอรู้จำภาษามือ [ข้อมูลการทำงานภายในโครงงาน]. ตรวจสอบเมื่อ 27 กันยายน 2569.','detector_performance_latest.json; สรุป 120 เฟรมล่าสุด จากทั้งหมด 627 เฟรม'),
('[5] คณะผู้พัฒนา HandVox. (2569). ผลการตรวจสอบส่วนข้อความ ประวัติ และการควบคุมการรับคำ [ชุดทดสอบและรายงานภายในโครงงาน]. ผล 26 กรณีตามที่รายงานในบทที่ 4 เมื่อ 27 กันยายน 2569.','tests/test_gesture_state.py; test_camera_sentence.py; test_sentence.py; test_history.py; test_detector_actions.py; test_detector_performance.py'),
('[6] คณะผู้พัฒนา HandVox. (2569, 23 กันยายน). ผลการประเมินผู้ทำท่าคนใหม่ person_03 สำหรับแบบจำลองรุ่น 20260915_061512_537476 [รายงานผลเบื้องต้นภายในโครงงาน].','experiments/20260915_061512_537476/external_evaluations/20260923_013726_285774_91def9df/report.md; metrics.json; ประเมิน 306 คลิป รวมคลิปรอตรวจคุณภาพ 228 คลิป'),
('[7] Bai, S., Kolter, J. Z., & Koltun, V. (2018). An Empirical Evaluation of Generic Convolutional and Recurrent Networks for Sequence Modeling. arXiv:1803.01271.','https://doi.org/10.48550/arXiv.1803.01271'),
]
for desc,path in entries:
    desc=re.sub(r'^\[\d+\] ', '', desc)
    p=para(refs,desc,'body');p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent=Cm(.75);p.paragraph_format.first_line_indent=Cm(-.75)
    p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(2)
    p.paragraph_format.line_spacing=Pt(18);p.paragraph_format.keep_with_next=True
    # Explicit opportunities at path separators prevent long technical paths from overflowing.
    path=path.replace('/','/\u200b').replace('_','_\u200b').replace(';',';\u200b')
    p=para(refs,'แหล่งข้อมูล: '+path,'body');p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent=Cm(.75);p.paragraph_format.first_line_indent=Cm(0)
    p.paragraph_format.keep_together=True;p.paragraph_format.line_spacing=Pt(16)
    p.paragraph_format.space_after=Pt(2)
    for r in p.runs:font(r,14)
refs.core_properties.title='เอกสารอ้างอิงประกอบบทที่ 5 โครงงาน HandVox'
refs.core_properties.author='คณะผู้จัดทำโครงงาน HandVox'
refs.save(OUT/'เอกสารอ้างอิง_บทที่ 5.docx')
print(json.dumps({'chapter_paragraphs':len(doc.paragraphs),'chapter_chars':sum(len(p.text.replace('\u200b','')) for p in doc.paragraphs),'references':len(entries),'output':'บทที่ 5.docx'},ensure_ascii=False))
